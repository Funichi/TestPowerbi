import datetime as dt
import json
import os
import shutil
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

HEAD = ["Seq", "Vertical", "Market", "Cliente", "Commessa", "SS", "Valore", "Oggetto", "Legal Entity", "tipo",
        "Data ins.", "Data Opt review", "Data consegna", "Solution D&A", "Stato", "Salesman", "Referenti vertical",
        "Tecnologie di riferimento D&A", "Note"]


def make_xlsx(path):
    from openpyxl.worksheet.table import Table
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "gare e offerte"
    ws.append(HEAD)
    d = dt.datetime
    data = [
        ["=ROW()-1", "Enterprise", "Energy&Utility", "Terna", "$EII4506", "TER1", 1, "Uno", None, "Gara", d(2026, 9, 1), None, d(2026, 10, 16), "Melloncelli / Scamuzzo", "WIP"],
        ["=ROW()-1", "Enterprise", "Energy&Utility", "TERNA ", None, "TER2", 1, "Due", None, "Gara", d(2026, 5, 1), None, d(2026, 5, 20), "Scamuzzo", "won"],
        ["=ROW()-1", "Finance", "Finance", "Generali", None, "GEN4", 1, "Tre", None, "RFI", d(2026, 8, 1), None, d(2026, 8, 10), "Grando", "Waiting For response"],
        ["=ROW()-1", "Finance", "Finance", "Intesa", None, None, 1, "Quattro", None, "RFI", None, None, None, None, "Closed"],
        ["=ROW()-1", "Public sector", " public  sector ", "Inps", None, "INP1", 1, "Cinque", None, "Gara", d(2026, 3, 2), None, d(2026, 3, 9), "Grando", "No bid"],
        ["=ROW()-1", "Enterprise", None, "Enel", None, "ENE1", 1, "Sei", None, "Offerta diretta", d(2025, 3, 2), None, d(2025, 3, 20), "Grando", "Lost"],
    ]
    for r in data:
        ws.append(r)
    from openpyxl.styles import PatternFill
    for r, rgb in ((2, "FFFFFF00"), (6, "FFFFC000"), (7, "FFFF0000")):                      # giallo, arancione, rosso sulla cella Cliente
        ws.cell(r, 4).fill = PatternFill("solid", start_color=rgb, end_color=rgb)
    ws.cell(4, 4).fill = PatternFill("solid", start_color="FFFFFFFF", end_color="FFFFFFFF")   # bianco: non è un livello
    ws.add_table(Table(displayName="Tabella4", ref=f"A1:S{len(data) + 1}"))
    dm = wb.create_sheet("Domini")
    dm.append([None, "Chiave", "Settore", "Codice", "Anno", None, "Vertical", None, "Market", None, "tipo", None, "Stato"])
    rows = [("Enterprise2026", "Enterprise", "$EII4506", 2026), ("Finance2026", "Finance", "$EII4505", 2026)]
    verts = ["Enterprise", "Public sector", "Finance"]
    mk = ["Energy&Utility", "Finance", "Public Sector"]
    tp = ["Gara", "RFI", "Offerta diretta"]
    st = ["Waiting For response", "NoBID", "Sospeso", "WIP", "Won", "Lost"]
    for i in range(6):
        r = rows[i] if i < len(rows) else (None, None, None, None)
        dm.append([None, r[0], r[1], r[2], r[3], None, verts[i] if i < 3 else None, None, mk[i] if i < 3 else None, None,
                   tp[i] if i < 3 else None, None, st[i]])
    wb.save(path)


@pytest.fixture()
def client(tmp_path, monkeypatch):
    x = tmp_path / "t.xlsx"
    make_xlsx(x)
    cfg = tmp_path / "config.json"
    cfg.write_text(json.dumps({"excel_path": str(x), "db_path": str(tmp_path / "monitoraggio.db"), "writer": "openpyxl"}))
    monkeypatch.setenv("GARE_CONFIG", str(cfg))
    import importlib
    import app as appmod
    importlib.reload(appmod)
    appmod.app.testing = True
    return appmod.app.test_client(), x, tmp_path


def sheet(x):
    return openpyxl.load_workbook(x)["gare e offerte"]


NEW = {"Cliente": "Terna", "Vertical": "enterprise", "Market": "Finance", "Oggetto": "Nuova", "tipo": "Gara",
       "Data ins.": "2026-10-10", "Data consegna": "2026-10-20", "Solution D&A": "Grando", "Stato": "WIP", "SS": "TER3", "Commessa": "$EII4506"}


def rows_by_cliente(c):
    return {r["Cliente"]: r for r in c.get("/api/data").get_json()["rows"]}


def test_primo_import_e_suggerimenti(client):
    c, x, tmp = client
    d = c.get("/api/data").get_json()                         # al primo avvio carica l'Excel nel database
    assert len(d["rows"]) == 6 and (tmp / "monitoraggio.db").exists()
    assert d["domini"]["vertical"] == ["Enterprise", "Public sector", "Finance"]
    st = c.get("/api/state").get_json()
    assert st["imported"] and st["pending"] == {"new": 0, "mod": 0}
    s = c.post("/api/suggest", json={"Cliente": "terna", "Vertical": "Enterprise", "Data consegna": "2026-11-01"}).get_json()
    assert s["ss"] == "TER3" and s["commessa"] == "$EII4506"
    assert c.post("/api/suggest", json={"Cliente": "Nuovo Cliente"}).get_json()["ss"] == "NUO1"


def test_modifiche_locali_non_toccano_l_excel(client):
    c, x, tmp = client
    before = sheet(x).max_row
    assert c.post("/api/rows", json=NEW).status_code == 200
    assert c.post("/api/rows", json=NEW).status_code == 409            # duplicato
    new = [r for r in c.get("/api/data").get_json()["rows"] if r["Oggetto"] == "Nuova"][0]
    assert new["_state"] == "new" and new["Vertical"] == "Enterprise" and new["Seq"] is None
    gen = rows_by_cliente(c)["Generali"]
    assert c.put(f"/api/rows/{gen['_row']}", json={"changes": {"Stato": "Won", "Note": "ok"}}).status_code == 200
    assert c.get("/api/state").get_json()["pending"] == {"new": 1, "mod": 1}
    assert sheet(x).max_row == before and sheet(x).cell(4, 15).value == "Waiting For response"      # Excel intatto
    # annulla modifiche e elimina la riga nuova
    assert c.post(f"/api/rows/{gen['_row']}/revert").status_code == 200
    assert c.delete(f"/api/rows/{new['_row']}").status_code == 200
    assert c.get("/api/state").get_json()["pending"] == {"new": 0, "mod": 0}
    assert c.delete(f"/api/rows/{gen['_row']}").status_code == 400    # le righe dell'Excel non si eliminano da qui
    assert c.post("/api/rows", json={**NEW, "Stato": "Boh", "Oggetto": "x"}).status_code == 400


def test_modificare_e_rimettere_il_valore_originale_non_e_una_modifica(client):
    c, x, _ = client
    gen = rows_by_cliente(c)["Generali"]
    c.put(f"/api/rows/{gen['_row']}", json={"changes": {"Stato": "Won"}})
    assert c.get("/api/state").get_json()["pending"]["mod"] == 1
    c.put(f"/api/rows/{gen['_row']}", json={"changes": {"Stato": "Waiting For response"}})
    assert c.get("/api/state").get_json()["pending"]["mod"] == 0


def test_export_accoda_e_aggiorna_poi_rilegge(client):
    c, x, tmp = client
    c.post("/api/rows", json=NEW)
    gen = rows_by_cliente(c)["Generali"]
    c.put(f"/api/rows/{gen['_row']}", json={"changes": {"Stato": "Won", "Note": "chiusa"}})
    p = c.get("/api/pending").get_json()
    assert len(p["nuove"]) == 1 and p["modificate"][0]["campi"][0]["campo"] in ("Stato", "Note")
    r = c.post("/api/export")
    assert r.status_code == 200, r.get_json()
    assert (r.get_json()["nuove"], r.get_json()["modificate"], r.get_json()["righe"]) == (1, 1, [8])
    ws = sheet(x)
    assert ws.tables["Tabella4"].ref == "A1:S8" and ws.cell(8, 4).value == "Terna" and ws.cell(8, 8).value == "Nuova"
    assert ws.cell(8, 1).value == "=ROW()-1" and ws.cell(8, 15).value == "WIP"
    assert ws.cell(4, 15).value == "Won" and ws.cell(4, 19).value == "chiusa"          # riga esistente aggiornata
    assert ws.cell(3, 4).value == "TERNA "                                           # le altre righe non si toccano
    assert len(list((tmp / "backup").glob("*.xlsx"))) == 1                          # backup dell'Excel prima di scrivere
    assert c.get("/api/state").get_json()["pending"] == {"new": 0, "mod": 0}          # database di nuovo uguale al file
    assert len(c.get("/api/data").get_json()["rows"]) == 7


def test_export_si_ferma_se_la_riga_nel_file_e_cambiata(client):
    c, x, _ = client
    gen = rows_by_cliente(c)["Generali"]
    c.put(f"/api/rows/{gen['_row']}", json={"changes": {"Stato": "Won"}})
    c.post("/api/rows", json=NEW)
    wb = openpyxl.load_workbook(x)
    wb["gare e offerte"].cell(4, 4).value = "Altro cliente"                          # qualcuno ha cambiato il file
    wb.save(x)
    r = c.post("/api/export")
    assert r.status_code == 409 and "cambiata" in r.get_json()["error"]
    ws = sheet(x)
    assert ws.max_row == 7 and ws.cell(4, 15).value == "Waiting For response"          # nessuna scrittura parziale


def test_aggiorna_dall_excel_chiede_conferma_e_scarta(client):
    c, x, tmp = client
    c.post("/api/rows", json=NEW)
    r = c.post("/api/import", json={})
    assert r.status_code == 409 and r.get_json()["conferma"]
    r = c.post("/api/import", json={"force": True})
    assert r.status_code == 200 and r.get_json()["backup"]
    assert len(list((tmp / "backup").glob("monitoraggio_*.db"))) == 1                # copia del database prima di scartare
    assert c.get("/api/state").get_json()["pending"] == {"new": 0, "mod": 0}
    assert len(c.get("/api/data").get_json()["rows"]) == 6
    # modifica fatta nel file Excel -> si vede dopo "Aggiorna"
    wb = openpyxl.load_workbook(x)
    wb["gare e offerte"].cell(4, 15).value = "Lost"
    wb.save(x)
    assert rows_by_cliente(c)["Generali"]["Stato"] == "Waiting For response"
    c.post("/api/import", json={})
    assert rows_by_cliente(c)["Generali"]["Stato"] == "Lost"


def test_pulizia_solo_suggerimenti(client):
    c, x, _ = client
    p = c.get("/api/clean/preview").get_json()
    kinds = {(i["kind"], i["new"]) for i in p}
    assert ("Stato uniformato", "Won") in kinds and ("Stato uniformato", "NoBID") in kinds
    assert not any(k[1] == "Sospeso" for k in kinds)          # Closed resta uno stato a sé
    assert ("Maiuscole/elenco", "Public Sector") in kinds
    assert all(i["xl"] for i in p) and any(i["cliente"] == "Generali" for i in p) is False
    assert c.post("/api/clean/apply", json={"items": p}).status_code in (404, 405)      # non scrive più nell'Excel
    assert sheet(x).cell(3, 15).value == "won"


def test_normalizzazione_in_lettura(client):
    c, x, _ = client
    rows = rows_by_cliente(c)
    assert rows["Terna"]["Stato"] == "WIP" and rows["TERNA"]["Stato"] == "Won"
    assert rows["Inps"]["Stato"] == "NoBID" and rows["Inps"]["Market"] == "Public Sector"
    assert rows["Intesa"]["Stato"] == "Closed"                  # stato a sé, non Sospeso
    assert rows["Enel"]["Market"] == "(non indicato)"
    summ = c.get("/api/clean/summary").get_json()
    assert {"campo": "Stato", "da": "No bid", "a": "NoBID", "righe": 1, "scrivibile": True} in summ
    assert {"campo": "Market", "da": "(vuoto)", "a": "(non indicato)", "righe": 1, "scrivibile": False} in summ
    assert sheet(x).cell(6, 15).value == "No bid"                # il file non è stato toccato dalla lettura


def test_cruscotto_kpi_trend_esito_eta(client, tmp_path):
    import datetime as dt
    from core.db import Store
    from core.logic import dashboard
    from core.reader import read_excel
    c, x, _ = client
    st = Store(str(tmp_path / "x.db"))
    st.import_all(read_excel(str(x)), str(x))
    d = st.load()
    db = dashboard(d["rows"], d["domini"], dt.date(2026, 10, 15), 12)
    k = db["kpi"]
    assert k["wip"] == 1 and k["waiting"] == 1
    assert k["inserite_ytd"] == {"cur": 4, "prev": 1} and k["won_ytd"] == {"cur": 1, "prev": 0}
    assert k["winrate12"]["won"] == 1 and k["winrate12"]["chiuse"] == 1      # Closed e NoBID esclusi
    t = db["trend"]
    assert t["mese"] == 10 and t["inserite"]["cur"][8] == 1 and t["inserite"]["prev"][2] == 1 and t["won"]["cur"][4] == 1
    tipi = {e["nome"]: e for e in dashboard(d["rows"], d["domini"], dt.date(2026, 10, 15), None)["esito_tipo"]}   # periodo: tutto
    assert tipi["Gara"]["totale"] == 3 and tipi["Gara"]["rate"] is None      # meno di 5 chiuse: nessun win rate
    assert tipi["RFI"]["Closed"] == 1 and tipi["Offerta diretta"]["Lost"] == 1
    assert tipi["Gara"]["NoBID"] == 1 and "Sospeso" not in tipi["Gara"]      # Sospeso e NoBID insieme come "No Bid"
    top = db["top_clienti"]
    assert top[0]["nome"] in ("Terna", "TERNA ", "TERNA") and top[0]["totale"] == 2 and top[0]["won"] == 1 and len(top[0]["rows"]) == 2
    assert db["pillole"]["attenzione"] == {"totale": 1, "rosse": 0, "arancioni": 0, "gialle": 1}
    assert db["attenzione_chart"]["wip"] == {"3": 0, "2": 0, "1": 1, "0": 0}                                   # Terna (WIP) è gialla
    assert db["attenzione_chart"]["tutte"] == {"3": 1, "2": 1, "1": 1, "0": 3}                                 # Enel rosso, Inps arancione
    wt = {w["nome"]: w for w in db["wip_tipo"]}                                                                 # oggi 15/10: la consegna del 16/10 è questa settimana
    assert wt["Gara"] == {"nome": "Gara", "scadute": 0, "settimana": 1, "prossima": 0, "oltre": 0, "senza_data": 0, "totale": 1}
    eta = {f["k"]: f["n"] for f in db["eta_attesa"]["fasce"]}
    assert eta["31-90"] == 1 and eta["nd"] == 0 and db["eta_attesa"]["totale"] == 1
    assert db["pillole"]["imminenti"] == 1 and db["pillole"]["scadute"] == 0


def test_attenzione_colore_dalla_cella_e_scrittura_nell_excel(client):
    from openpyxl.styles import PatternFill  # noqa: F401
    c, x, _ = client
    rows = rows_by_cliente(c)
    assert (rows["Terna"]["Attenzione"], rows["Inps"]["Attenzione"], rows["Enel"]["Attenzione"], rows["TERNA"]["Attenzione"]) == (1, 2, 3, 0)
    # il colore si cambia nell'app (database locale) e finisce nell'Excel solo con l'esportazione
    assert c.put(f"/api/rows/{rows['Generali']['_row']}", json={"changes": {"Attenzione": 3}}).status_code == 200
    assert c.put(f"/api/rows/{rows['Terna']['_row']}", json={"changes": {"Attenzione": 0}}).status_code == 200
    assert c.put(f"/api/rows/{rows['Terna']['_row']}", json={"changes": {"Attenzione": 7}}).status_code == 400
    assert c.get("/api/state").get_json()["pending"] == {"new": 0, "mod": 2}
    assert sheet(x).cell(4, 4).fill.fgColor.rgb == "FFFFFFFF"                      # Excel intatto
    c.post("/api/rows", json={**NEW, "Attenzione": 2})
    p = c.get("/api/pending").get_json()
    assert any(f["campo"] == "Attenzione" and f["dopo"].startswith("Urgente") for m in p["modificate"] for f in m["campi"])
    assert p["nuove"][0]["Attenzione"].startswith("Da attenzionare")
    assert c.post("/api/export").status_code == 200
    ws = sheet(x)
    assert ws.cell(4, 4).fill.fgColor.rgb == "00FF0000" or ws.cell(4, 4).fill.fgColor.rgb.endswith("FF0000")     # Generali: rosso
    assert ws.cell(2, 4).fill.fill_type is None                                                               # Terna: colore tolto
    assert ws.cell(8, 4).fill.fgColor.rgb.endswith("FFC000")                                                  # riga nuova: arancione
    assert ws.cell(6, 4).fill.fgColor.rgb.endswith("FFC000") and ws.cell(7, 4).fill.fgColor.rgb.endswith("FF0000")   # le altre righe non si toccano
    allrows = c.get("/api/data").get_json()["rows"]                                                           # dopo l'export il database coincide col file
    att = {r["Oggetto"]: r["Attenzione"] for r in allrows}
    assert (att["Tre"], att["Uno"], att["Cinque"], att["Nuova"]) == (3, 0, 2, 2)


def test_annullare_modifiche_ripristina_anche_il_colore(client):
    c, x, _ = client
    t = rows_by_cliente(c)["Terna"]
    c.put(f"/api/rows/{t['_row']}", json={"changes": {"Attenzione": 3}})
    assert c.get("/api/state").get_json()["pending"]["mod"] == 1
    assert c.post(f"/api/rows/{t['_row']}/revert").status_code == 200
    assert rows_by_cliente(c)["Terna"]["Attenzione"] == 1
    c.put(f"/api/rows/{t['_row']}", json={"changes": {"Attenzione": 3}})
    c.put(f"/api/rows/{t['_row']}", json={"changes": {"Attenzione": 1}})                     # rimesso com'era: non è più una modifica
    assert c.get("/api/state").get_json()["pending"]["mod"] == 0


def test_nuova_offerta_senza_data_consegna_va_a_fine_mese(client):
    c, x, _ = client
    body = {k: v for k, v in NEW.items() if k != "Data consegna"}
    assert c.post("/api/rows", json={**body, "Data ins.": "2026-02-10", "Oggetto": "feb"}).status_code == 200
    assert c.post("/api/rows", json={**body, "Data ins.": "2026-10-16", "Oggetto": "ott"}).status_code == 200
    assert c.post("/api/rows", json={**NEW, "Oggetto": "con data"}).status_code == 200                 # se c'è, non si tocca
    got = {r["Oggetto"]: r["Data consegna"] for r in c.get("/api/data").get_json()["rows"] if r["Oggetto"] in ("feb", "ott", "con data")}
    assert got == {"feb": "2026-02-28", "ott": "2026-10-31", "con data": "2026-10-20"}


def test_ss_progressivo_per_anno_e_duplicati_solo_nello_stesso_anno(client):
    c, x, _ = client
    # Terna ha TER1 e TER2 nel 2026: nel 2026 il prossimo è TER3, nel 2027 si riparte da TER1
    assert c.post("/api/suggest", json={"Cliente": "Terna", "Data ins.": "2026-10-10"}).get_json()["ss"] == "TER3"
    assert c.post("/api/suggest", json={"Cliente": "Terna", "Data ins.": "2027-01-10"}).get_json()["ss"] == "TER1"
    # stesso codice in anni diversi: non è un duplicato; nello stesso anno sì
    c.post("/api/rows", json={**NEW, "SS": "TER2", "Data ins.": "2025-10-10", "Data consegna": "2025-10-20", "Oggetto": "Vecchia"})
    titoli = lambda: [a["titolo"] for a in c.get("/api/dashboard").get_json()["attenzioni"]]
    assert not any("SS duplicato" in t for t in titoli())
    c.post("/api/rows", json={**NEW, "SS": "TER2", "Oggetto": "Stesso anno"})
    assert any("SS duplicato nello stesso anno" in t for t in titoli())


def test_scheda_cliente_tempi_e_cartella(client, tmp_path):
    c, x, _ = client
    s = c.get("/api/cliente", query_string={"nome": "terna"}).get_json()
    assert s["kpi"]["totale"] == 2 and s["kpi"]["won"] == 1 and s["kpi"]["wip"] == 1
    assert s["kpi"]["tempi"] == {"n": 1, "media": 19.0, "mediana": 19.0}          # la WIP non conta: solo la Won (1/5 -> 20/5)
    assert s["cronologia"][0]["Oggetto"] == "Uno" and [a["anno"] for a in s["anni"]] == [2026]
    assert s["media"]["tempi"]["n"] == 3 and c.get("/api/cliente", query_string={"nome": "boh"}).status_code == 400
    assert s["cartella"]["stato"] == "nessuna"
    t = c.get("/api/tempi").get_json()
    assert t["totale"]["n"] == 3 and {g["nome"] for g in t["tipo"]} == {"Gara", "RFI", "Offerta diretta"} and t["solution"]
    # cartelle: una uguale al nome, una solo simile (si sceglie e si ricorda)
    base = tmp_path / "Clienti"
    for n in ("Terna", "Generali Assicurazioni"):
        (base / n).mkdir(parents=True)
    cfg = json.loads((tmp_path / "config.json").read_text())
    cfg["clienti_dir"] = str(base)
    (tmp_path / "config.json").write_text(json.dumps(cfg))
    import app as appmod
    opened = []
    appmod.open_folder = opened.append
    assert c.get("/api/cliente", query_string={"nome": "Terna"}).get_json()["cartella"]["stato"] == "ok"
    g = c.get("/api/cliente", query_string={"nome": "Generali"}).get_json()["cartella"]
    assert g["stato"] == "scegli" and g["candidati"] == ["Generali Assicurazioni"]
    assert c.post("/api/cliente/cartella", json={"cliente": "Generali"}).status_code == 404
    assert c.post("/api/cliente/cartella", json={"cliente": "Generali", "cartella": "../x"}).status_code == 404
    assert c.post("/api/cliente/cartella", json={"cliente": "Generali", "cartella": "Generali Assicurazioni"}).status_code == 200
    assert c.get("/api/cliente", query_string={"nome": "Generali"}).get_json()["cartella"] == {"stato": "ok", "nome": "Generali Assicurazioni", "candidati": []}
    assert opened and opened[-1].endswith("Generali Assicurazioni")


def test_varianti_cliente_e_unione_nel_database_locale(client):
    c, x, _ = client
    c.post("/api/rows", json={**NEW, "Cliente": "Terna S.p.A.", "Oggetto": "Variante"})
    v = c.get("/api/clienti/varianti").get_json()
    assert len(v) == 1 and {n["nome"] for n in v[0]["nomi"]} == {"Terna", "TERNA", "Terna S.p.A."} and v[0]["suggerito"] in ("Terna", "TERNA")
    r = c.post("/api/clienti/unisci", json={"nomi": ["Terna", "TERNA", "Terna S.p.A."], "in": "Terna"}).get_json()
    assert r["righe"] == 2 and c.get("/api/clienti/varianti").get_json() == []
    assert c.get("/api/state").get_json()["pending"]["mod"] == 1 and sheet(x).cell(3, 4).value == "TERNA "      # solo database locale


def test_scheda_cliente_filtrata_per_anno(client):
    c, x, _ = client
    c.post("/api/rows", json={**NEW, "Cliente": "Enel", "Oggetto": "Nuova Enel", "SS": "ENE2", "Data ins.": "2026-10-10", "Data consegna": "2026-10-20", "Stato": "Won"})
    tutti = c.get("/api/cliente", query_string={"nome": "Enel"}).get_json()
    assert tutti["kpi"]["totale"] == 2 and tutti["anni_disponibili"] == [2025, 2026]
    a25 = c.get("/api/cliente", query_string={"nome": "Enel", "anno": 2025}).get_json()
    assert a25["kpi"]["totale"] == 1 and a25["kpi"]["lost"] == 1 and a25["kpi"]["won"] == 0 and a25["anni_disponibili"] == [2025, 2026]
    assert c.get("/api/cliente", query_string={"nome": "Terna", "anno": 2025}).get_json()["kpi"]["totale"] == 0      # nessuna offerta in quell'anno: scheda vuota, non errore


def test_cruscotto_per_anno_solare(client):
    c, x, _ = client
    d26 = c.get("/api/dashboard", query_string={"anno": 2026}).get_json()
    d25 = c.get("/api/dashboard", query_string={"anno": 2025}).get_json()
    tot = lambda d: sum(g["totale"] for g in d["esito_tipo"])
    assert tot(d26) == 4 and tot(d25) == 1                       # Enel (2025) solo nel 2025; la riga senza Data ins. non entra in un anno
    assert [t["nome"] for t in d25["top_clienti"]] == ["Enel"]
    assert d25["kpi"]["wip"] == 1                                # i KPI "in corso" restano quelli di oggi
