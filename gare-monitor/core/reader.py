"""Lettura dell'Excel (sola lettura, su una copia temporanea: il file originale non viene mai bloccato)."""
import datetime as dt
import os
import re
import shutil
import tempfile
import warnings

import openpyxl

from .normalize import normalize_rows

SHEET_MAIN = "gare e offerte"
SHEET_DOMINI = "Domini"

DATE_FIELDS = ["Data ins.", "Data Opt review", "Data consegna"]
# Campi che l'app mostra/modifica. Le colonne economiche e "Valore" restano intatte nel file.
FIELDS = [
    "Vertical", "Market", "Cliente", "Commessa", "SS", "Oggetto", "Legal Entity", "tipo",
    "Data ins.", "Data Opt review", "Data consegna", "Solution D&A", "Stato", "Salesman",
    "Referenti vertical", "Tecnologie di riferimento D&A", "Note",
]


ISO_RE = re.compile(r"^\d{4}-\d{2}-\d{2}([T ]\d{2}:\d{2}(:\d{2})?)?$")


def to_date(v):
    """datetime/date/stringa gg/mm/aaaa o aaaa-mm-gg -> date, altrimenti None."""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        s = v.strip()
        if ISO_RE.match(s):
            s = s[:10]
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y"):
            try:
                return dt.datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return None


def to_iso(v):
    d = to_date(v)
    return d.isoformat() if d else None


def clean_text(v):
    if v is None:
        return None
    s = str(v).replace("\xa0", " ")
    s = re.sub(r"[ \t]+", " ", s).strip()
    return s or None


def row_sig(cliente, oggetto, data_ins):
    """Impronta della riga: serve a verificare, prima di scrivere, che la riga non sia cambiata nel frattempo."""
    c = (clean_text(cliente) or "").casefold()
    o = (clean_text(oggetto) or "").casefold()
    return f"{c}|{o}|{to_iso(data_ins) or ''}"


def _column_list(ws, header):
    """Valori sotto l'intestazione `header` (riga 1) fino alla prima cella vuota."""
    col = None
    for c in ws[1]:
        if clean_text(c.value) == header:
            col = c.column
            break
    out = []
    if col is None:
        return out
    r = 2
    while ws.cell(r, col).value not in (None, ""):
        out.append(clean_text(ws.cell(r, col).value))
        r += 1
    return out


# Livello di attenzione = colore di riempimento della cella del nome cliente, come si fa oggi nell'Excel
ATT_COLORS = {1: "FFFF00", 2: "FFC000", 3: "FF0000"}          # giallo = rilevante, arancione = da attenzionare, rosso = urgente
ATT_LABELS = {0: "nessuna", 1: "Rilevante (giallo)", 2: "Da attenzionare (arancione)", 3: "Urgente (rosso)"}


def fill_level(cell):
    """0 se la cella non ha uno dei tre colori; altrimenti 1, 2 o 3."""
    f = cell.fill
    if f is None or f.fill_type != "solid":
        return 0
    c = f.fgColor
    if c is None or c.type != "rgb" or not isinstance(c.rgb, str) or len(c.rgb) < 6:
        return 0
    rgb = c.rgb[-6:].upper()
    for lvl, ref in ATT_COLORS.items():
        if sum(abs(int(rgb[i:i + 2], 16) - int(ref[i:i + 2], 16)) for i in (0, 2, 4)) <= 50:
            return lvl
    return 0


def _cell_value(v):
    """Valore di cella pronto per il database: date come testo ISO, il resto invariato."""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if (v.hour, v.minute, v.second) == (0, 0, 0) else v.isoformat(timespec="seconds")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat(timespec="seconds")
    return v


def read_excel(path):
    """Legge TUTTE le colonne del foglio principale (+ le liste di Domini), su una copia temporanea del file.
    Ritorna {'headers': [...], 'rows': [{'xl': n_riga, 'cells': {colonna: valore}}], 'domini': {...}}."""
    tmp = tempfile.mkdtemp(prefix="gare_")
    try:
        copy = os.path.join(tmp, "copia.xlsx")
        try:
            shutil.copyfile(path, copy)
        except OSError as e:
            if os.name != "nt":
                raise
            try:
                from .snapshot import snapshot_via_excel
                snapshot_via_excel(path, copy)
            except Exception as e2:
                raise PermissionError(
                    f"Non riesco a leggere il file ({e}). Provato anche tramite Excel: {e2}. "
                    "Controlla che il file sia scaricato da OneDrive ('Mantieni sempre su questo dispositivo') e non bloccato.") from e2
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(copy, data_only=True)
        if SHEET_MAIN not in wb.sheetnames:
            raise ValueError(f"Nel file manca il foglio '{SHEET_MAIN}'")
        ws = wb[SHEET_MAIN]
        headers, cols = [], {}
        for c in ws[1]:
            h = clean_text(c.value)
            if not h:
                continue
            base, n = h, 2
            while h in cols:                       # intestazioni duplicate: " (2)", " (3)"…
                h = f"{base} ({n})"
                n += 1
            cols[h] = c.column
            headers.append(h)
        missing = [f for f in ("Cliente", "Stato", "Data ins.") if f not in cols]
        if missing:
            raise ValueError(f"Nel foglio '{SHEET_MAIN}' mancano le colonne: {', '.join(missing)}")
        rows = []
        for r in range(2, ws.max_row + 1):
            cells = {h: _cell_value(ws.cell(r, cols[h]).value) for h in headers}
            if not (clean_text(cells.get("Cliente")) or clean_text(cells.get("Oggetto"))):
                continue
            rows.append({"xl": r, "cells": cells, "att": fill_level(ws.cell(r, cols["Cliente"]))})
        domini = _read_domini(wb[SHEET_DOMINI] if SHEET_DOMINI in wb.sheetnames else None)
        return {"headers": headers, "rows": rows, "domini": domini}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def build_row(rid, xl, cells, orig=None, state="orig", att=0):
    """Riga per l'app a partire dai valori del database.
    cells = valori attuali; orig = valori importati dall'Excel (None per le righe nuove)."""
    row = {"_row": rid, "_xl": xl, "_state": state, "Attenzione": att or 0}
    seq = cells.get("Seq")
    row["Seq"] = seq if isinstance(seq, (int, float)) and not isinstance(seq, bool) else (xl - 1 if xl else None)
    invalid, textdates = [], []
    for f in FIELDS:
        v = cells.get(f)
        if f in DATE_FIELDS:
            iso = to_iso(v)
            row[f] = iso
            if v not in (None, "") and iso is None:
                invalid.append(f)
            if isinstance(v, str) and iso and not ISO_RE.match(v.strip()):
                textdates.append(f)
        else:
            row[f] = clean_text(v)
    row["_invalid_dates"] = invalid
    if textdates:
        row["_textdates"] = textdates
    base = orig if orig is not None else cells
    row["_raw"] = {k: (str(base[k]) if base.get(k) is not None else None) for k in ("Stato", "Market") if k in base}
    row["_sig"] = row_sig(cells.get("Cliente"), cells.get("Oggetto"), cells.get("Data ins."))
    return row


def _read_domini(ws):
    d = {"vertical": [], "market": [], "tipo": [], "stato": [], "commesse": {}}
    if ws is None:
        return d
    d["vertical"] = _column_list(ws, "Vertical")
    d["market"] = _column_list(ws, "Market")
    d["tipo"] = _column_list(ws, "tipo")
    seen = []
    for s in _column_list(ws, "Stato"):
        if s not in seen:
            seen.append(s)
    d["stato"] = seen
    # tabella Settore/Codice/Anno -> commessa
    cols = {}
    for c in ws[1]:
        h = clean_text(c.value)
        if h in ("Settore", "Codice", "Anno"):
            cols[h] = c.column
    if len(cols) == 3:
        r = 2
        while ws.cell(r, cols["Settore"]).value not in (None, ""):
            sett, cod, anno = (ws.cell(r, cols[k]).value for k in ("Settore", "Codice", "Anno"))
            if isinstance(sett, str) and cod and anno:
                d["commesse"][f"{sett.strip().casefold()}|{int(anno)}"] = str(cod).strip()
            r += 1
    return d
