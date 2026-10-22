"""Referenze: esportazione in Excel (anche modello per i caricamenti massivi), importazione con anteprima e importazione di schede JSON con unione."""
import difflib
import io
import re
import secrets
import time
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import referenze_cfg as cfg
from .referenze import MODIFICABILI, RefError, cliente_key, mancanti_automatici, normalizza, valida

SEP = " ; "
PUNTO = "• "
# (campo, intestazione, larghezza): stesso ordine dei campi dell'archivio
COLONNE = [("id", "ID", 11), ("origine", "Origine", 16), ("contatore", "Contatore", 10), ("cliente", "Cliente", 26), ("titolo", "Titolo", 44),
           ("titolo_excel", "Titolo (Excel originale)", 32), ("ambito", "Ambito", 24), ("tipologie", "Tipologie", 38), ("mercato", "Mercato", 24),
           ("segmento", "Segmento", 18), ("anno", "Anno", 8), ("temi", "Temi", 28), ("tecnologie", "Tecnologie", 30), ("pitch", "Racconto 30 secondi", 60),
           ("esteso", "Racconto esteso", 90), ("fonti", "Fonti", 30), ("documenti", "Documenti", 16), ("confidenza", "Confidenza", 12), ("materiale", "Materiale", 12),
           ("campi_mancanti", "Campi mancanti", 32), ("da_verificare", "Da verificare", 11), ("note", "Note", 44), ("descrizione_excel", "Descrizione (Excel)", 50),
           ("tecnologie_excel", "Tecnologie (Excel)", 28), ("referente_aidata", "Referente AI&Data", 22), ("referente_mercato", "Referente mercato", 22),
           ("referente_progetto", "Referente progetto", 22), ("owner", "Owner", 16), ("archiviata", "Archiviata", 11)]
LISTE = ("tipologie", "temi", "tecnologie", "fonti", "documenti", "campi_mancanti")
BOOL = ("da_verificare", "archiviata")
NON_MODIFICABILI = {"id", "origine", "contatore"}
DA_CONFRONTARE = [c for c, _, _ in COLONNE if c not in NON_MODIFICABILI]
# nomi alternativi delle intestazioni (si confrontano senza maiuscole, accenti e punteggiatura)
EXTRA_ALIAS = {"id": ["idreferenza", "codice", "ref"], "contatore": ["nriga", "riga", "n"], "titolo_excel": ["titoloexcel", "titolooriginale", "titolomaster"],
               "tipologie": ["tipologia"], "pitch": ["pitch", "in30secondi", "racconto30secondi", "racconto30s"], "esteso": ["raccontoesteso", "esteso"],
               "descrizione_excel": ["descrizione", "descrizioneexcel"], "tecnologie_excel": ["tecnologieexcel"], "referente_aidata": ["referenteaidata", "referenteaiedata", "referenteai"],
               "campi_mancanti": ["mancante", "campimancanti"], "da_verificare": ["daverificare", "verificare"], "owner": ["proprietario"]}
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _n(s):
    return re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().casefold())


ALIAS = {}
for _c, _h, _w in COLONNE:
    ALIAS[_n(_h)] = _c
    ALIAS[_n(_c)] = _c
for _c, _l in EXTRA_ALIAS.items():
    for _a in _l:
        ALIAS.setdefault(_n(_a), _c)


# ---------------------------------------------------------------- formato delle celle (racconto esteso e liste)
def esteso_a_testo(esteso):
    parti = []
    for s in esteso or []:
        t = s.get("testo")
        corpo = "\n".join(PUNTO + str(x) for x in t) if isinstance(t, list) else str(t or "")
        parti.append(f"## {s['sezione']}\n{corpo}")
    return "\n\n".join(parti)


def testo_a_esteso(testo):
    """Inverso di esteso_a_testo: '## Sezione' apre una sezione; se tutte le righe iniziano con '• ' è un elenco, altrimenti un paragrafo."""
    sezioni, cur = [], None
    for riga in str(testo or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if riga.startswith("## "):
            cur = {"sezione": riga[3:].strip(), "righe": []}
            sezioni.append(cur)
        elif cur is not None:
            cur["righe"].append(riga)
        elif riga.strip():
            raise RefError("Il racconto esteso deve iniziare con una riga '## Nome sezione'")
    out = []
    for s in sezioni:
        piene = [r for r in s["righe"] if r.strip()]
        if piene and all(r.lstrip().startswith("•") for r in piene):
            out.append({"sezione": s["sezione"], "testo": [r.lstrip()[1:].strip() for r in piene]})
        else:
            out.append({"sezione": s["sezione"], "testo": "\n".join(s["righe"]).strip()})
    return out


def _liste(v):
    return [x.strip() for x in str(v or "").split(";") if x.strip()]


def _bool(v):
    return str(v or "").strip().casefold() in ("sì", "si", "s", "true", "vero", "1", "x", "yes", "y")


def _pulisci(s):
    return _CTRL.sub("", s) if isinstance(s, str) else s


# ---------------------------------------------------------------- forma canonica (per confrontare file e archivio senza falsi cambiamenti)
def canon(campo, v):
    if campo in LISTE:
        return [str(x).strip() for x in (v or []) if str(x).strip()]
    if campo == "esteso":
        out = []
        for s in v or []:
            t = s.get("testo")
            out.append({"sezione": str(s.get("sezione", "")).strip(), "testo": [str(x).strip() for x in t if str(x).strip()] if isinstance(t, list) else str(t or "").strip()})
        return out
    if campo in BOOL:
        return bool(v)
    return str(v if v is not None else "").strip()


# ---------------------------------------------------------------- esportazione
def esporta(store):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Referenze"
    ws.append([h for _, h, _ in COLONNE])
    for r in store.elenco(archiviate=True):
        riga = []
        for campo, _, _ in COLONNE:
            v = r.get(campo)
            if campo in LISTE:
                v = SEP.join(v or [])
            elif campo == "esteso":
                v = esteso_a_testo(v)
            elif campo in BOOL:
                v = "Sì" if v else "No"
            elif campo == "contatore":
                v = v if isinstance(v, int) else None
            riga.append(_pulisci(v) if v != "" else None)
        ws.append(riga)
    testata = PatternFill("solid", start_color="FFC51B88", end_color="FFC51B88")
    for i, (campo, _, w) in enumerate(COLONNE, 1):
        c = ws.cell(1, i)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFFFF"), testata, Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, i)
            cell.alignment = Alignment(vertical="top", wrap_text=campo in ("pitch", "esteso", "note", "descrizione_excel", "titolo", "temi", "tecnologie", "tipologie", "fonti", "campi_mancanti"))
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"                                   # un testo che inizia con "=" non deve diventare una formula
            if campo == "anno":
                cell.number_format = "@"
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    # Domini
    dm = wb.create_sheet("Domini")
    colonne = [("Ambito", cfg.AMBITI), ("Tipologie AI", cfg.TIPOLOGIE["AI"]), ("Tipologie Data Management & Reporting", cfg.TIPOLOGIE["Data Management & Reporting"]),
               ("Mercato", cfg.MERCATI), ("Campi mancanti", cfg.CAMPI_MANCANTI), ("Materiale", cfg.MATERIALE), ("Confidenza", cfg.CONFIDENZE), ("Sezioni standard del racconto esteso", cfg.SEZIONI_STANDARD)]
    for i, (titolo, valori) in enumerate(colonne, 1):
        dm.cell(1, i, titolo).font = Font(bold=True)
        dm.column_dimensions[dm.cell(1, i).column_letter].width = 42
        for j, v in enumerate(valori, 2):
            dm.cell(j, i, v)
    # elenchi a tendina per le colonne con valori ammessi (aiutano i caricamenti massivi)
    for campo, col, n in (("ambito", "A", len(cfg.AMBITI)), ("mercato", "D", len(cfg.MERCATI)), ("materiale", "F", len(cfg.MATERIALE)), ("confidenza", "G", len(cfg.CONFIDENZE))):
        idx = [c for c, _, _ in COLONNE].index(campo) + 1
        dv = DataValidation(type="list", formula1=f"=Domini!${col}$2:${col}${n + 1}", allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{ws.cell(2, idx).coordinate}:{ws.cell(max(ws.max_row, 500), idx).coordinate}")
    # Leggimi
    lg = wb.create_sheet("Leggimi")
    righe = [
        "Referenze — file di esportazione e modello per i caricamenti massivi",
        "",
        "Ogni riga del foglio «Referenze» è una referenza. Le intestazioni delle colonne non vanno cambiate.",
        "• Riga con ID esistente (es. REF-012): aggiorna quella referenza. Si leggono solo le colonne presenti nel file.",
        "• Riga senza ID: crea una nuova referenza (l'ID lo assegna l'app).",
        "• Prima di applicare l'app mostra un'anteprima: nuove, aggiornate, righe con errori e i campi che cambiano.",
        "",
        "Formato delle celle",
        f"• Liste (Tipologie, Temi, Tecnologie, Fonti, Documenti, Campi mancanti): voci separate da «{SEP.strip()}».",
        "• Racconto esteso: ogni sezione inizia con una riga «## Nome sezione»; gli elenchi puntati con «• » a inizio riga; altrimenti è un paragrafo.",
        "• Da verificare e Archiviata: Sì / No.",
        "• Anno: 4 cifre (es. 2024) oppure vuoto.",
        "• Ambito, Mercato, Tipologie, Materiale, Confidenza, Campi mancanti: solo valori del foglio «Domini».",
        "• ID, Origine e Contatore non si modificano da qui.",
        "",
        "Campi obbligatori per una nuova referenza: Cliente, Titolo, Ambito, Mercato.",
    ]
    for i, t in enumerate(righe, 1):
        lg.cell(i, 1, t).font = Font(bold=(i == 1 or t == "Formato delle celle"), size=13 if i == 1 else 11)
    lg.column_dimensions["A"].width = 130
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- lettura del file
def leggi_xlsx(fileobj):
    """File Excel -> (righe, colonne_presenti, avvisi). righe = [{'n': numero riga Excel, 'dati': {campo: valore}}]."""
    try:
        wb = openpyxl.load_workbook(fileobj, data_only=True, read_only=True)
    except Exception as e:
        raise RefError(f"Non riesco ad aprire il file Excel: {e}") from e
    ws = wb["Referenze"] if "Referenze" in wb.sheetnames else wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    intestazione = next(it, None)
    if not intestazione:
        raise RefError("Il foglio è vuoto")
    mappa, avvisi = {}, []
    for i, h in enumerate(intestazione):
        if h is None or str(h).strip() == "":
            continue
        campo = ALIAS.get(_n(h))
        if campo and campo not in mappa.values():
            mappa[i] = campo
        else:
            avvisi.append(f"Colonna ignorata: «{h}»")
    if not mappa:
        raise RefError("Non riconosco nessuna colonna: usa il file esportato dall'app come modello")
    righe = []
    for n, row in enumerate(it, 2):
        dati = {}
        for i, campo in mappa.items():
            v = row[i] if i < len(row) else None
            if campo == "contatore":
                dati[campo] = int(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
            elif campo in LISTE:
                dati[campo] = _liste(v)
            elif campo in BOOL:
                dati[campo] = _bool(v)
            elif campo == "esteso":
                dati[campo] = testo_a_esteso(_pulisci(v))
            elif campo == "anno":
                dati[campo] = str(int(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer() else str(v or "").strip()
            else:
                dati[campo] = str(v).strip() if v is not None else ""
        if any(dati.get(c) not in ("", [], None, False) for c in dati if c not in ("id",)) or dati.get("id"):
            righe.append({"n": n, "dati": dati})
    wb.close()
    return righe, set(mappa.values()), avvisi


# ---------------------------------------------------------------- anteprima e applicazione (Excel)
_PENDENTI = {}


def _pulisci_pendenti():
    for t in [t for t, v in _PENDENTI.items() if time.time() - v["ts"] > 1800]:
        _PENDENTI.pop(t, None)


def anteprima_excel(store, fileobj):
    righe, presenti, avvisi = leggi_xlsx(fileobj)
    esistenti = {r["id"]: r for r in store.elenco(archiviate=True)}
    nuove, aggiornate, invariate, errori, viste = [], [], 0, [], set()
    for riga in righe:
        n, dati = riga["n"], riga["dati"]
        rid = (dati.get("id") or "").strip()
        try:
            if rid:
                if rid in viste:
                    raise RefError(f"ID duplicato nel file: {rid}")
                viste.add(rid)
                cur = esistenti.get(rid)
                if not cur:
                    raise RefError(f"ID {rid} non presente nell'archivio (lascia l'ID vuoto per crearla come nuova)")
                cambi = {}
                for campo in DA_CONFRONTARE:
                    if campo not in presenti or campo not in dati:
                        continue
                    if campo in ("materiale", "confidenza") and not dati[campo]:
                        continue                                              # cella vuota = non cambiare
                    if canon(campo, dati[campo]) != canon(campo, cur[campo]):
                        cambi[campo] = dati[campo]
                if not cambi:
                    invariate += 1
                    continue
                valida(normalizza({**cur, **{k: v for k, v in cambi.items() if k != "archiviata"}}), cur)
                aggiornate.append({"n": n, "id": rid, "cliente": cur["cliente"], "titolo": cur["titolo"], "cambi": cambi,
                                   "campi": [{"campo": c, "etichetta": dict((k, h) for k, h, _ in COLONNE)[c], "prima": cur[c], "dopo": v} for c, v in cambi.items()]})
            else:
                nuova = {k: v for k, v in dati.items() if k in MODIFICABILI or k == "archiviata"}
                valida(normalizza(nuova))
                nuove.append({"n": n, "cliente": nuova.get("cliente", ""), "titolo": nuova.get("titolo", ""), "dati": nuova})
        except RefError as e:
            errori.append({"n": n, "id": rid, "cliente": dati.get("cliente", ""), "titolo": dati.get("titolo", ""), "errore": str(e)})
    _pulisci_pendenti()
    token = secrets.token_hex(12)
    _PENDENTI[token] = {"ts": time.time(), "nuove": nuove, "aggiornate": aggiornate}
    return {"token": token, "righe": len(righe), "invariate": invariate, "avvisi": avvisi, "errori": errori,
            "nuove": [{k: v for k, v in x.items() if k != "dati"} for x in nuove],
            "aggiornate": [{k: v for k, v in x.items() if k != "cambi"} for x in aggiornate]}


def applica_excel(store, token):
    p = _PENDENTI.pop(token, None)
    if not p:
        raise RefError("L'anteprima è scaduta: ricarica il file")
    store.backup()                                                           # copia di sicurezza prima di scrivere
    creati = aggiornati = 0
    errori = []
    for x in p["nuove"]:
        try:
            d = dict(x["dati"])
            arch = d.pop("archiviata", False)
            nuova = store.crea(d, origine="Import Excel")
            if arch:
                store.archivia(nuova["id"], True)
            creati += 1
        except RefError as e:
            errori.append({"n": x["n"], "errore": str(e)})
    for x in p["aggiornate"]:
        try:
            cambi = dict(x["cambi"])
            arch = cambi.pop("archiviata", None)
            if cambi:
                store.aggiorna(x["id"], cambi)
            if arch is not None:
                store.archivia(x["id"], arch)
            aggiornati += 1
        except RefError as e:
            errori.append({"n": x["n"], "errore": str(e)})
    return {"creati": creati, "aggiornati": aggiornati, "errori": errori}


# ---------------------------------------------------------------- importa scheda (JSON) con unione
def _norm_titolo(t):
    return re.sub(r"[^a-z0-9 ]", "", unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().casefold())


def schede_da_json(payload):
    if isinstance(payload, dict):
        payload = payload.get("referenze", [payload])
    if not isinstance(payload, list) or not payload or not all(isinstance(x, dict) for x in payload):
        raise RefError("Mi aspetto una scheda, un elenco di schede, oppure {\"referenze\": [...]}")
    return payload


def anteprima_schede(store, payload):
    esistenti = store.elenco(archiviate=True)
    out = []
    for i, grezza in enumerate(schede_da_json(payload)):
        s = normalizza({k: v for k, v in grezza.items() if k in MODIFICABILI or k in ("id", "origine")})
        errori = []
        try:
            valida(s)
        except RefError as e:
            errori = str(e).split("; ")
        cand = []
        for r in esistenti:
            punti = 0.5 * (cliente_key(r["cliente"]) == cliente_key(s["cliente"]) and bool(s["cliente"])) + 0.5 * difflib.SequenceMatcher(None, _norm_titolo(s["titolo"]), _norm_titolo(r["titolo"])).ratio()
            if s["id"] and s["id"] == r["id"]:
                punti = 1.0
            if punti >= 0.45:
                cand.append({"id": r["id"], "cliente": r["cliente"], "titolo": r["titolo"], "punteggio": round(punti, 2), "archiviata": r["archiviata"]})
        cand.sort(key=lambda c: -c["punteggio"])
        out.append({"indice": i, "scheda": s, "errori": errori, "candidati": cand[:5], "suggerimento": cand[0]["id"] if cand and cand[0]["punteggio"] >= 0.8 else "nuova"})
    return {"schede": out}


def applica_schede(store, azioni):
    store.backup()
    creati = uniti = saltati = 0
    errori = []
    for a in azioni or []:
        tipo = a.get("tipo")
        try:
            if tipo == "nuova":
                d = {k: v for k, v in (a.get("scheda") or {}).items() if k in MODIFICABILI}
                store.crea(d, origine=(a.get("scheda") or {}).get("origine") or "Import scheda")
                creati += 1
            elif tipo == "unisci":
                store.aggiorna(a["id"], {k: v for k, v in (a.get("campi") or {}).items() if k in MODIFICABILI})
                uniti += 1
            else:
                saltati += 1
        except (RefError, KeyError) as e:
            errori.append({"indice": a.get("indice"), "errore": str(e)})
    return {"creati": creati, "uniti": uniti, "saltati": saltati, "errori": errori}
